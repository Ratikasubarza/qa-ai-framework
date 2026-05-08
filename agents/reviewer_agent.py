"""
Test Case Reviewer Agent — reviews test cases against PRD.
Input:  PRD content (text/file/Jira) + existing test cases (CSV/JSON)
Output: Coverage gap analysis + traceability matrix + findings + recommendations
"""
import json
import csv
import structlog
from pathlib import Path
from dataclasses import dataclass
from core.llm_client import llm_client, ModelTier
from core.circuit_breaker import get_breaker
from core import input_sanitizer
from core.prompt_manager import prompt_manager
from skills.reporting.excel_writer import ExcelWriter
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = structlog.get_logger(__name__)

AGENT_NAME     = "reviewer_agent"
PROMPT_NAME    = "reviewer"
PROMPT_VERSION = "1.0.0"


@dataclass
class ReviewReport:
    total_requirements:       int
    total_test_cases:         int
    coverage_percentage:      float
    findings_critical:        int
    findings_major:           int
    findings_minor:           int
    traceability_matrix:      list
    detailed_findings:        list
    missing_test_cases:       list
    prd_clarifications:       list
    overall_quality_score:    float
    recommendation:           str
    raw_data:                 dict


class ReviewerAgent:
    def __init__(self):
        self.breaker = get_breaker(AGENT_NAME)

    def review(
        self,
        prd_content:   str,
        test_cases_path: str | None = None,
        test_cases_json: list | None = None,
        feature:       str = "Feature",
    ) -> ReviewReport:
        """
        Review test cases against PRD.
        Accepts test cases from CSV file path OR JSON list.
        """
        logger.info("reviewer_agent_start", feature=feature)

        # ── Load test cases ───────────────────────────────────
        if test_cases_path:
            test_cases_text = self._load_csv(test_cases_path)
        elif test_cases_json:
            test_cases_text = json.dumps(test_cases_json, indent=2)
        else:
            test_cases_text = "No test cases provided — perform gap analysis from PRD only."

        # ── Sanitize inputs ───────────────────────────────────
        clean_prd = input_sanitizer.sanitize(prd_content, source="prd")
        clean_tc  = input_sanitizer.sanitize(test_cases_text, source="test_cases")

        # ── Format prompt ─────────────────────────────────────
        template = prompt_manager.get(PROMPT_NAME)
        system   = template.format(
            prd_content=clean_prd.content,
            test_cases=clean_tc.content,
        )

        # ── LLM call ──────────────────────────────────────────
        def _call():
            return llm_client.complete_text(
                system=system,
                user=f"Review test cases for: {feature}",
                tier=ModelTier.COMPLEX,
                agent_name=AGENT_NAME,
                json_mode=True,
            )

        response = self.breaker.call(_call)
        data     = self._parse_response(response.content)

        report = ReviewReport(
            total_requirements=data.get("summary", {}).get("total_requirements", 0),
            total_test_cases=data.get("summary", {}).get("total_test_cases_reviewed", 0),
            coverage_percentage=data.get("summary", {}).get("coverage_percentage", 0),
            findings_critical=data.get("summary", {}).get("findings_critical", 0),
            findings_major=data.get("summary", {}).get("findings_major", 0),
            findings_minor=data.get("summary", {}).get("findings_minor", 0),
            traceability_matrix=data.get("traceability_matrix", []),
            detailed_findings=data.get("detailed_findings", []),
            missing_test_cases=data.get("missing_test_cases", []),
            prd_clarifications=data.get("prd_clarifications_needed", []),
            overall_quality_score=data.get("overall_quality_score", 0),
            recommendation=data.get("recommendation", ""),
            raw_data=data,
        )

        logger.info(
            "review_complete",
            coverage=report.coverage_percentage,
            critical=report.findings_critical,
            missing=len(report.missing_test_cases),
        )
        return report

    def export_report(self, report: ReviewReport, feature: str) -> dict[str, Path]:
        """Export review report to Excel and CSV."""
        from datetime import datetime
        date_str     = datetime.now().strftime("%Y%m%d")
        feature_slug = feature.lower().replace(" ", "_")
        output_dir   = Path("reports/excel")
        output_dir.mkdir(parents=True, exist_ok=True)

        # ── Excel report ──────────────────────────────────────
        excel_path = output_dir / f"review_{feature_slug}_{date_str}.xlsx"
        self._write_excel_report(report, feature, excel_path)

        # ── CSV findings ──────────────────────────────────────
        csv_path = output_dir / f"review_{feature_slug}_{date_str}.csv"
        self._write_csv_report(report, csv_path)

        logger.info("review_exported", excel=str(excel_path), csv=str(csv_path))
        return {"excel": excel_path, "csv": csv_path}

    def _write_excel_report(self, report: ReviewReport, feature: str, path: Path):
        wb = Workbook()

        # ── Sheet 1: Summary ──────────────────────────────────
        ws = wb.active
        ws.title = "Summary"
        dark_fill   = PatternFill("solid", fgColor="1E293B")
        green_fill  = PatternFill("solid", fgColor="CCFFCC")
        red_fill    = PatternFill("solid", fgColor="FFCCCC")
        yellow_fill = PatternFill("solid", fgColor="FFE5CC")
        bold_white  = Font(bold=True, color="FFFFFF", size=12)
        bold_dark   = Font(bold=True, color="1E293B", size=11)
        thin = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        summary_data = [
            ("Feature",                  feature),
            ("Total Requirements",        report.total_requirements),
            ("Total Test Cases Reviewed", report.total_test_cases),
            ("Coverage %",               f"{report.coverage_percentage:.0f}%"),
            ("Quality Score",            f"{report.overall_quality_score}/10"),
            ("Critical Findings",         report.findings_critical),
            ("Major Findings",            report.findings_major),
            ("Minor Findings",            report.findings_minor),
            ("Missing Test Cases",        len(report.missing_test_cases)),
            ("Recommendation",            report.recommendation),
        ]
        for row_idx, (label, value) in enumerate(summary_data, 1):
            c1 = ws.cell(row=row_idx, column=1, value=label)
            c2 = ws.cell(row=row_idx, column=2, value=str(value))
            c1.fill, c1.font, c1.border = dark_fill, bold_white, thin
            c2.border = thin
            c2.alignment = Alignment(wrap_text=True)
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 60

        # ── Sheet 2: Traceability Matrix ──────────────────────
        ws2 = wb.create_sheet("Traceability")
        headers = ["Requirement ID", "Requirement Summary", "Covering Test Cases", "Status"]
        for col, h in enumerate(headers, 1):
            cell = ws2.cell(row=1, column=col, value=h)
            cell.fill, cell.font, cell.border = dark_fill, bold_white, thin

        status_colors = {"covered": "CCFFCC", "partial": "FFE5CC", "uncovered": "FFCCCC"}
        for row_idx, item in enumerate(report.traceability_matrix, 2):
            status = item.get("status", "uncovered")
            fill   = PatternFill("solid", fgColor=status_colors.get(status, "FFFFFF"))
            row_data = [
                item.get("requirement_id", ""),
                item.get("requirement_summary", ""),
                ", ".join(item.get("covering_test_cases", [])),
                status.upper(),
            ]
            for col_idx, val in enumerate(row_data, 1):
                cell = ws2.cell(row=row_idx, column=col_idx, value=val)
                cell.fill, cell.border = fill, thin
        for i, w in enumerate([15, 50, 30, 12], 1):
            ws2.column_dimensions[get_column_letter(i)].width = w

        # ── Sheet 3: Findings ─────────────────────────────────
        ws3 = wb.create_sheet("Findings")
        headers3 = ["Severity", "Affected Test Case", "Issue", "Suggestion"]
        for col, h in enumerate(headers3, 1):
            cell = ws3.cell(row=1, column=col, value=h)
            cell.fill, cell.font, cell.border = dark_fill, bold_white, thin

        sev_colors = {"Critical": "FFCCCC", "Major": "FFE5CC", "Minor": "FFFFCC"}
        for row_idx, finding in enumerate(report.detailed_findings, 2):
            sev   = finding.get("severity", "Minor")
            fill  = PatternFill("solid", fgColor=sev_colors.get(sev, "FFFFFF"))
            row_data = [
                sev,
                finding.get("affected_test_case", ""),
                finding.get("issue", ""),
                finding.get("suggestion", ""),
            ]
            for col_idx, val in enumerate(row_data, 1):
                cell = ws3.cell(row=row_idx, column=col_idx, value=val)
                cell.fill  = fill
                cell.border = thin
                cell.alignment = Alignment(wrap_text=True)
        for i, w in enumerate([12, 25, 50, 50], 1):
            ws3.column_dimensions[get_column_letter(i)].width = w
            ws3.row_dimensions[1].height = 25

        # ── Sheet 4: Missing Test Cases ───────────────────────
        ws4 = wb.create_sheet("Missing Test Cases")
        headers4 = ["Priority", "Title", "Reason", "Covers Requirement"]
        for col, h in enumerate(headers4, 1):
            cell = ws4.cell(row=1, column=col, value=h)
            cell.fill, cell.font, cell.border = dark_fill, bold_white, thin

        priority_colors = {"P0": "FFCCCC", "P1": "FFE5CC", "P2": "CCFFCC"}
        for row_idx, tc in enumerate(report.missing_test_cases, 2):
            pri  = tc.get("priority", "P1")
            fill = PatternFill("solid", fgColor=priority_colors.get(pri, "FFFFFF"))
            row_data = [
                pri,
                tc.get("title", ""),
                tc.get("reason", ""),
                tc.get("covers_requirement", ""),
            ]
            for col_idx, val in enumerate(row_data, 1):
                cell = ws4.cell(row=row_idx, column=col_idx, value=val)
                cell.fill  = fill
                cell.border = thin
                cell.alignment = Alignment(wrap_text=True)
        for i, w in enumerate([10, 40, 50, 30], 1):
            ws4.column_dimensions[get_column_letter(i)].width = w

        wb.save(path)

    def _write_csv_report(self, report: ReviewReport, path: Path):
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Severity", "Affected Test Case", "Issue", "Suggestion"])
            for finding in report.detailed_findings:
                writer.writerow([
                    finding.get("severity", ""),
                    finding.get("affected_test_case", ""),
                    finding.get("issue", ""),
                    finding.get("suggestion", ""),
                ])

    def _load_csv(self, path: str) -> str:
        """Load CSV test cases into text for prompt."""
        try:
            with open(path, "r", encoding="utf-8") as f:
                return f.read()
        except FileNotFoundError:
            logger.warning("csv_not_found", path=path)
            return "Test cases file not found."

    def _parse_response(self, raw: str) -> dict:
        raw = raw.strip().lstrip("```json").lstrip("```").rstrip("```").strip()
        try:
            return json.loads(raw)
        except json.JSONDecodeError as e:
            logger.error("reviewer_parse_failed", error=str(e))
            return {
                "summary": {"total_requirements": 0, "total_test_cases_reviewed": 0,
                            "coverage_percentage": 0, "findings_critical": 0,
                            "findings_major": 0, "findings_minor": 0},
                "traceability_matrix": [],
                "detailed_findings": [],
                "missing_test_cases": [],
                "prd_clarifications_needed": [],
                "overall_quality_score": 0,
                "recommendation": "Parse error — retry",
            }


reviewer_agent = ReviewerAgent()
